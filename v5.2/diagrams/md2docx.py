"""MD → DOCX（图片页横向）。用法：python3 md2docx.py in.md out.docx（图片路径相对 md 所在目录）"""
import re, sys, os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
md, out = sys.argv[1], sys.argv[2]; base = os.path.dirname(os.path.abspath(md))
doc=Document(); st=doc.styles['Normal']; st.font.name='Microsoft YaHei'; st.font.size=Pt(10); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
def runs(p,t):
    for j,part in enumerate(re.split(r'\*\*',t)):
        r=p.add_run(part.replace('`','')); r.bold=j%2==1
def orient(sec, land):
    sec.orientation = WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
    w,h = Cm(29.7), Cm(21.0)
    sec.page_width, sec.page_height = (w,h) if land else (h,w)
    for a in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(sec,a,Cm(1.5))
orient(doc.sections[0], False)
lines=open(md,encoding='utf-8').read().split('\n'); i=0
while i<len(lines):
    l=lines[i]; m=re.match(r'!\[(.*)\]\((.*)\)',l)
    if m:
        orient(doc.add_section(WD_SECTION.NEW_PAGE), True)
        doc.add_picture(os.path.join(base, m.group(2)), width=Cm(26.5)); runs(doc.add_paragraph(), "图："+m.group(1))
        orient(doc.add_section(WD_SECTION.NEW_PAGE), False); i+=1; continue
    if l.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            c=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in c if x): rows.append(c)
            i+=1
        t=doc.add_table(rows=len(rows),cols=max(map(len,rows))); t.style='Table Grid'
        for a,r in enumerate(rows):
            for b,x in enumerate(r):
                cell=t.cell(a,b); cell.text=''; runs(cell.paragraphs[0],x)
                for rr in cell.paragraphs[0].runs: rr.font.size=Pt(8.5); rr.bold=rr.bold or a==0
        continue
    h=re.match(r'(#+) (.*)',l)
    if h: doc.add_heading(h.group(2),min(len(h.group(1)),3))
    elif l.startswith('> '): runs(doc.add_paragraph(),l[2:])
    elif re.match(r'\s*- ',l): runs(doc.add_paragraph(style='List Bullet'),re.sub(r'^\s*- ','',l))
    elif re.match(r'\d+\. ',l): runs(doc.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',l))
    elif l.strip(): runs(doc.add_paragraph(),l)
    i+=1
doc.save(out); print(out)
