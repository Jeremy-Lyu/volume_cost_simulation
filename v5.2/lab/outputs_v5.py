"""v5 输出升级：成本瀑布、影子价格（固定模式后有限差分）、一页管理层决策单（MD+PNG）、统一 Excel"""
import os as _root_os, sys as _root_sys
_root_sys.path.insert(0, _root_os.path.dirname(_root_os.path.dirname(_root_os.path.abspath(__file__))))
import copy, os
import engine as en
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
import cjkfont; cjkfont.use()
def waterfall(ra, rb, path, la="A", lb="B"):
    keys = [k for k in dict.fromkeys(list(ra["cost"]) + list(rb["cost"])) if not k.startswith(("权重:", "−")) and k not in ("物料", "缺口惩罚(流失)")]
    diff = [(k, (rb["cost"].get(k, 0) - ra["cost"].get(k, 0)) / 1e4) for k in keys]
    diff = [x for x in diff if abs(x[1]) >= 0.5]; diff.sort(key=lambda x: -abs(x[1]))
    a0 = ra["kpi"]["运营成本"] / 1e4; b0 = rb["kpi"]["运营成本"] / 1e4
    fig, ax = plt.subplots(figsize=(15, 7)); cur = a0
    ax.bar(0, a0, color="#546e7a"); ax.text(0, a0, f"{a0:,.0f}", ha="center", va="bottom", fontsize=11)
    for i, (k, v) in enumerate(diff, 1):
        ax.bar(i, v, bottom=cur if v >= 0 else cur + v, color="#c62828" if v > 0 else "#2e7d32")
        ax.text(i, max(cur, cur + v), f"{v:+,.0f}", ha="center", va="bottom", fontsize=10); cur += v
    n = len(diff) + 1; ax.bar(n, b0, color="#1565c0"); ax.text(n, b0, f"{b0:,.0f}", ha="center", va="bottom", fontsize=11)
    ax.set_xticks(range(n + 1)); ax.set_xticklabels([la] + [k for k, _ in diff] + [lb], rotation=30, ha="right", fontsize=11)
    lo = min(a0, b0) * 0.97; ax.set_ylim(lo, max(a0, b0, cur) * 1.01)
    ax.set_ylabel("万元"); ax.set_title(f"成本瀑布：{la} → {lb}（按科目，示例数据；红=增加 绿=减少）", fontsize=14)
    fig.tight_layout(); fig.savefig(path, dpi=120); plt.close(fig)
    return [dict(科目=k, 差额_万元=round(v, 1)) for k, v in diff]
def shadow_prices(base, r0, procs=6):
    """固定基准方案的模式 y，放宽一个资源，看运营成本下降多少（元/单位）。LP 可行域内的边际价值，⚠ 只对小幅放宽有效"""
    y = r0["y"]; d0 = copy.deepcopy(base); d0["g"]["solver_time_limit"] = 30
    ref = en.solve_case(copy.deepcopy(d0), "影子基准", fix_y=y)
    specs = []
    for s in base["shops"]:
        e = copy.deepcopy(d0); e["shop"][s]["ot_max_hours_month"] += 1; specs.append((f"{s} 常规加班上限 +1h/人", 1, e))
        if s != base["shops"][-1]:
            e = copy.deepcopy(d0); e["shop"][s]["buffer_cap"] += 100; specs.append((f"{s} 缓存上限 +100台", 100, e))
        e = copy.deepcopy(d0); e["cap"] = {k: (v * 1.01 if k[0] == s else v) for k, v in e["cap"].items()}; specs.append((f"{s} 车型产能上限 +1%", 1, e))
    e = copy.deepcopy(d0); e["g"]["ot_hard_cap_per_head_month"] += 1; specs.append(("加班绝对上限 +1h/人 ⚠HR", 1, e))
    res = en.run_many([(e, nm, dict(fix_y=y)) for nm, _, e in specs], procs); rows = []
    for (nm, unit, _), r in zip(specs, res):
        dv = (ref["kpi"]["运营成本"] - r["kpi"]["运营成本"]) / unit
        rows.append(dict(资源=nm, 每单位节省_元=round(dv), 全年节省_万元=round(dv * unit / 1e4, 1), 说明="0 表示该资源当前不紧"))
    return sorted(rows, key=lambda x: -x["每单位节省_元"])
def check_baseline(ctx, tol_wan=1):
    """v5.1：决策单口径自检——表中"成本优先"行必须与结论文字/瀑布起点为同一基准（同一次求解），否则报错"""
    base_row = next((r for r in ctx["方案"] if r["方案"] == "成本优先"), None)
    if base_row is None: return []
    err = []
    if ctx.get("瀑布基准_万元") is not None and abs(ctx["瀑布基准_万元"] - base_row["运营成本_万元"]) > tol_wan:
        err.append(f"瀑布起点 {ctx['瀑布基准_万元']} ≠ 表中成本优先 {base_row['运营成本_万元']}")
    if ctx.get("结论基准_万元") is not None and abs(ctx["结论基准_万元"] - base_row["运营成本_万元"]) > tol_wan:
        err.append(f"结论基准 {ctx['结论基准_万元']} ≠ 表中成本优先 {base_row['运营成本_万元']}")
    return err
def decision_sheet(ctx, md_path, png_path):
    err = check_baseline(ctx); assert not err, f"决策单基准混用：{err}"
    L = [f"# {ctx['标题']}", "", f"> {ctx['版本']}　｜　示例数据　｜　⚠ = 暂定参数", "", f"**基准口径**：{ctx.get('基准口径', '成本优先')}", "", f"**结论**：{ctx['结论']}", "",
         "## 1. 备选方案对比", "", "| 方案 | 运营成本(万元) | 单车(元) | 人数最大月波动 | 延期(台·月) | 流失(台) | 超限加班(h) |", "|:--|--:|--:|--:|--:|--:|--:|"]
    for r in ctx["方案"]: L.append(f"| {r['方案']} | {r['运营成本_万元']:,} | {r['单车_元']:,} | {r['人数最大月波动']:.0%} | {r['延期台月']} | {r['流失台']} | {r['超限加班h']} |")
    L += ["", "## 2. 为什么（成本差 Top 科目）", ""] + [f"- {x['科目']}：{x['差额_万元']:+,.0f} 万元" for x in ctx["瀑布"][:5]]
    L += ["", "## 3. 风险（预测分布下 40 个样本）", ""] + [f"- {r['方案']}：平均 {r['平均成本_万元']:,} 万，P90 {r['P90']:,} 万，P(流失) {r['P_流失']}，P(成本超预算) {r['P_成本超预算']}" for r in ctx["风险"]]
    L += ["", "## 4. 安全边界", ""] + [f"- {r['维度']}：{r['临界值']}" for r in ctx["边界"]]
    L += ["", "## 5. 最值钱的资源（影子价格）", ""] + [f"- {r['资源']}：每单位省 {r['每单位节省_元']:,} 元" for r in ctx["影子"][:4]]
    L += ["", "## 6. 需要拍板", ""] + [f"- {x}" for x in ctx["拍板"]] + ["", "## 7. 假设与待确认 ⚠", ""] + [f"- {x}" for x in ctx["假设"]]
    open(md_path, "w", encoding="utf-8").write("\n".join(L))
    import textwrap
    txt = [w for x in L if not x.startswith("|:--") for w in (textwrap.wrap(x.replace("**", "").replace("# ", "").replace("> ", ""), 58) or [""])]
    h = (len(txt) + 1) // 2; k = next((j for j in range(h, len(txt)) if txt[j].startswith("#")), h)
    fig = plt.figure(figsize=(16.5, 11.7))
    fig.text(0.02, 0.98, "\n".join(txt[:k]), va="top", fontsize=10.5, linespacing=1.45); fig.text(0.51, 0.98, "\n".join(txt[k:]), va="top", fontsize=10.5, linespacing=1.45)
    fig.savefig(png_path, dpi=110); plt.close(fig)
def excel(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        if not rows: ws.append(["（无）"]); continue
        if isinstance(rows[0], str):
            for x in rows: ws.append([x])
            ws.column_dimensions["A"].width = 140; continue
        keys = list(dict.fromkeys(k for r in rows for k in r)); ws.append(keys)
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="37474F"); c.alignment = Alignment(wrap_text=True)
        for r in rows: ws.append([r.get(k, "") if not isinstance(r.get(k), (list, dict, tuple)) else str(r.get(k)) for k in keys])
        for i, k in enumerate(keys, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, min(40, len(str(k)) * 2 + 2))
        ws.freeze_panes = "A2"
    wb.save(path)
