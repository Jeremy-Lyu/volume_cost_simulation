import sys; sys.path.insert(0,'.')
from content_v52 import FIGS
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
IMG='/workspace/review/图表-v5.2/'
FONT='Noto Sans CJK SC'
d=Document()
sec=d.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Cm(29.7); sec.page_height=Cm(21.0)
for a in ('left_margin','right_margin'): setattr(sec,a,Cm(1.6))
sec.top_margin=Cm(1.3); sec.bottom_margin=Cm(1.2); sec.header_distance=Cm(0.6); sec.footer_distance=Cm(0.5)
def setfont(style,size=None,bold=None,color=None):
    style.font.name=FONT
    rpr=style.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for k in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(k),FONT)
    for k in ('w:asciiTheme','w:hAnsiTheme','w:eastAsiaTheme','w:cstheme'):
        if rf.get(qn(k)) is not None: del rf.attrib[qn(k)]
    if size: style.font.size=Pt(size)
    if bold is not None: style.font.bold=bold
    if color: style.font.color.rgb=RGBColor.from_string(color)
setfont(d.styles['Normal'],10)
d.styles['Normal'].paragraph_format.space_after=Pt(2)
h1=d.styles['Heading 1']; setfont(h1,15,True,'1F3864'); h1.paragraph_format.space_before=Pt(0); h1.paragraph_format.space_after=Pt(4); h1.paragraph_format.keep_with_next=True
for n in ('Title','Subtitle','TOC Heading'):
    try: setfont(d.styles[n])
    except KeyError: pass
# footer page number
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
def field(par,instr):
    r=par.add_run(); f1=OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'),'begin'); r._r.append(f1)
    r=par.add_run(); it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=instr; r._r.append(it)
    r=par.add_run(); f2=OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'),'separate'); r._r.append(f2)
    r=par.add_run('（右键更新域以生成目录）' if 'TOC' in instr else '1')
    r=par.add_run(); f3=OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'),'end'); r._r.append(f3)
r=fp.add_run('产量成本模拟器 · 重要图表汇编 v5.2 · 示例数据　第 '); r.font.size=Pt(8); r.font.color.rgb=RGBColor(0x88,0x88,0x88)
field(fp,'PAGE'); r=fp.add_run(' 页'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(0x88,0x88,0x88)
def para(text='',size=10,bold=False,color=None,align=None,after=2,before=0):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before)
    if align: p.alignment=align
    if text:
        r=p.add_run(text); r.font.size=Pt(size); r.bold=bold
        if color: r.font.color.rgb=RGBColor.from_string(color)
    return p
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill); tcPr.append(sh)
def borders(tbl,color='8EA9DB'):
    tblPr=tbl._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement(f'w:{e}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),color); b.append(x)
    tblPr.append(b)
def cellmargins(tbl,top=40,bot=40,lr=80):
    tblPr=tbl._tbl.tblPr; m=OxmlElement('w:tblCellMar')
    for k,v in (('top',top),('bottom',bot),('left',lr),('right',lr)):
        x=OxmlElement(f'w:{k}'); x.set(qn('w:w'),str(v)); x.set(qn('w:type'),'dxa'); m.append(x)
    tblPr.append(m)
def fixw(tbl,widths):
    tblPr=tbl._tbl.tblPr; lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    for j,gc in enumerate(tbl._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths[j].twips if hasattr(widths[j],'twips') else widths[j]/635)))
    for row in tbl.rows:
        for j,c in enumerate(row.cells): c.width=widths[j]
def celltext(cell,text,size=9.5,bold=False,color=None):
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(text); r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=RGBColor.from_string(color)
# ---------- cover ----------
para(after=0).add_run().add_break()
para('产量成本模拟器 · 重要图表汇编（v5.2 版）',28,True,'1F3864',WD_ALIGN_PARAGRAPH.CENTER,after=8,before=60)
para('架构与逻辑 · v4 基准与场景结果 · v5 决策层 · 合作与推广 · v5.1 能源与按天排产 · v5.2 人工弹性与换档',14,False,'404040',WD_ALIGN_PARAGRAPH.CENTER,after=6)
para('配套：《产量成本模型-要素拆解-v5.2.md》《工厂运营理论与计算手册-v5.2》《v5.2升级说明.md》',11,False,'595959',WD_ALIGN_PARAGRAPH.CENTER,after=4)
para('2026-10-18',11,False,'595959',WD_ALIGN_PARAGRAPH.CENTER,after=30)
t=d.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t,'C00000'); cellmargins(t,100,100,160)
c=t.rows[0].cells[0]; shade(c,'FFF4F4')
celltext(c,'数据声明：',10.5,True,'C00000')
c.paragraphs[0].add_run('本汇编中的全部数字均为模拟器的示例数据，不是沃尔沃真实数据，仅用于说明方法。标 ⚠ 的是暂定参数或待确认事项。v5.2 在 v5.1 基础上加入按天人工弹性、跨车间借调与 JPH 换档成本，基准年总成本 44,288 万元（v5.1 为 44,746 万元）；图 9、图 10–11、v5 决策层各图按 v5.2 重新生成；第五部分为 v5.1 能源与按天排产（v5.1 结果）；v5.2 新增图见第六部分').font.size=Pt(10.5)
fixw(t,[Cm(22)])
para(after=4,before=18)
para('阅读方式：每张图单独一页，依次为图题 → 图 → 信息表（来源文件、对应 MD 章节、怎么读、结论 / 要点、注意事项）。原图 PNG 均在 /workspace/review/图表-v5.2/ 目录下。',10.5,False,'404040',WD_ALIGN_PARAGRAPH.CENTER)
p=para(); p.add_run().add_break(WD_BREAK.PAGE)
# ---------- 目录 + 总览 ----------
para('目录',16,True,'1F3864',after=4)
p=para(); field(p,'TOC \\o "1-1" \\h \\z \\u')
p=para(); p.add_run().add_break(WD_BREAK.PAGE)
para('图表总览',16,True,'1F3864',after=4)
t=d.add_table(rows=1,cols=5); borders(t); cellmargins(t,20,20,60)
for j,h in enumerate(['图号','部分','图题','来源 / 生成方式','对应 MD 章节']): celltext(t.rows[0].cells[j],h,9,True); shade(t.rows[0].cells[j],'D9E2F3')
for i,f in enumerate(FIGS,1):
    r=t.add_row().cells
    src=f[3].split('（')[0]; how='重新生成' if '重新生成' in f[3] else ('新绘/重排' if ('新绘' in f[3] or '重新排版' in f[3]) else ('沿用原图' if '沿用' in f[3] else 'v4/v5 交付物'))
    for j,v in enumerate([f'图{i}',f[0],f[2],f'{src}｜{how}',f[4].replace('《产量成本模型-要素拆解-v5.md》','要素拆解 ').split('；')[0]]): celltext(r[j],v,8.5)
fixw(t,[Cm(1.3),Cm(3.2),Cm(8.0),Cm(7.0),Cm(6.9)])
# ---------- figures ----------
LBL=[('来源文件',3),('对应 MD 章节',4),('怎么读',5),('结论 / 要点',6),('注意 ⚠',7)]
for i,f in enumerate(FIGS,1):
    hp=d.add_paragraph(f'图{i}　{f[2]}',style='Heading 1'); hp.paragraph_format.page_break_before=True
    q=para(f'{f[0]}',8.5,False,'7F7F7F',after=3)
    w,h=Image.open(IMG+f[1]).size; ar=h/w
    rows=[(l,f[k]) for l,k in LBL if f[k]]
    tall=ar>0.72
    if tall:
        maxw,maxh=Cm(17.5),Cm(16.3)
        iw=min(maxw,int(maxh/ar))
        t=d.add_table(rows=1,cols=2); cellmargins(t,0,0,60)
        c0,c1=t.rows[0].cells
        c0.paragraphs[0].add_run().add_picture(IMG+f[1],width=iw); c0.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        it=c1.add_table(rows=0,cols=2); borders(it); cellmargins(it,30,30,70)
        for l,v in rows:
            rc=it.add_row().cells; celltext(rc[0],l,9,True); shade(rc[0],'D9E2F3'); celltext(rc[1],v,9, color=('C00000' if l.startswith('注意') else None))
        tw=Cm(26.5)-iw-Cm(0.4); fixw(it,[Cm(2.0),tw-Cm(2.0)])
        if not c1.paragraphs[0].text: 
            pe=c1.paragraphs[0]._p; pe.getparent().remove(pe)
        fixw(t,[iw+Cm(0.4),Cm(26.5)-iw-Cm(0.4)])
    else:
        # estimate text height to size image
        tl=sum(max(1,int(len(v)/88)+1) for _,v in rows); txt_cm=0.48*tl+0.25*len(rows)
        maxh=Cm(max(7.0,min(12.2,16.4-txt_cm))); maxw=Cm(26.0)
        iw=min(maxw,int(maxh/ar))
        pp=para(after=3,align=WD_ALIGN_PARAGRAPH.CENTER); pp.add_run().add_picture(IMG+f[1],width=iw)
        it=d.add_table(rows=0,cols=2); borders(it); cellmargins(it,25,25,70)
        for l,v in rows:
            rc=it.add_row().cells; celltext(rc[0],l,9,True); shade(rc[0],'D9E2F3'); celltext(rc[1],v,9, color=('C00000' if l.startswith('注意') else None))
        fixw(it,[Cm(2.4),Cm(24.1)])
d.save('/workspace/review/重要图表汇编-v5.2.docx'); print('saved')
