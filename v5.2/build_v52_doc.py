"""v5.2升级说明.md → v5.2升级说明.docx（python-docx；表格、图片、标题、列表）"""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
H = os.path.dirname(os.path.abspath(__file__))
def run_fmt(p, text):
    for i, seg in enumerate(re.split(r"\*\*", text)):
        r = p.add_run(seg.replace("`", "")); r.bold = bool(i % 2)
def main(md=os.path.join(os.path.dirname(H), "发布", "版本记录", "v5.2升级说明.md"), out=os.path.join(os.path.dirname(H), "发布", "版本记录", "v5.2升级说明.docx")):
    doc = Document(); st = doc.styles["Normal"]; st.font.name = "Noto Sans CJK SC"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    for nm in ("Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        hs = doc.styles[nm]; hs.font.name = "Noto Sans CJK SC"
        rpr = hs.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        if rf is None: rf = rpr.makeelement(qn("w:rFonts"), {}); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia"): rf.set(qn(a), "Noto Sans CJK SC")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"): rf.attrib.pop(qn(a), None)
    for s in doc.sections: s.left_margin = s.right_margin = Cm(2)
    lines = open(md, encoding="utf-8").read().split("\n"); i = 0
    while i < len(lines):
        L = lines[i]
        if L.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[:\- |]+\|$", lines[i]): rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.cell(r, c); cell.text = ""; p = cell.paragraphs[0]; run_fmt(p, val)
                    for rr in p.runs: rr.font.size = Pt(8.5); rr.bold = rr.bold or r == 0
            doc.add_paragraph(); continue
        m = re.match(r"!\[(.*)\]\((.*)\)", L)
        if m:
            p = os.path.join(H, m.group(2))
            if os.path.exists(p): doc.add_picture(p, width=Cm(16.5)); cap = doc.add_paragraph(m.group(1)); cap.runs[0].font.size = Pt(9)
        elif L.startswith("# "): doc.add_heading(L[2:], 0)
        elif L.startswith("## "): doc.add_heading(L[3:], 1)
        elif L.startswith("### "): doc.add_heading(L[4:], 2)
        elif L.startswith("> "): p = doc.add_paragraph(); run_fmt(p, L[2:]); [setattr(r.font, "color", None) for r in []]; p.runs[0].italic = True
        elif re.match(r"^\d+\. ", L): p = doc.add_paragraph(style="List Number"); run_fmt(p, re.sub(r"^\d+\. ", "", L))
        elif L.startswith("- "): p = doc.add_paragraph(style="List Bullet"); run_fmt(p, L[2:])
        elif L.strip(): p = doc.add_paragraph(); run_fmt(p, L)
        i += 1
    doc.save(out); print("已生成", out)
if __name__ == "__main__": main()
